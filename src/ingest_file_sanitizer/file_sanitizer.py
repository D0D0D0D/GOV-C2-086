# PART: ingest-file-sanitizer v0.1.1 (parts@1565fd9)
"""Magic-byte allowlist sanitizer for ingest file payloads."""

from __future__ import annotations

import csv
import io
import string
import zipfile
from dataclasses import dataclass, field
from typing import Any


_ZIP_MAGIC = b"PK\x03\x04"
_CSV_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class TextQualityPolicy:
    min_text_chars: int = 1
    min_printable_ratio: float = 0.85
    max_replacement_ratio: float = 0.01
    max_control_ratio: float = 0.02
    min_signal_ratio: float = 0.15


@dataclass
class FileSanitizeResult:
    channel: str | None = None
    accepted: bool = True
    reject_reason: str | None = None
    sanitized_text: str | None = None
    sanitized_rows: list[dict[str, Any]] = field(default_factory=list)
    rejected: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def sniff_channel(raw_bytes: bytes, declared_name: str = "", declared_content_type: str = "") -> str | None:
    """Sniff content. Declared name/content-type are intentionally ignored."""

    if not raw_bytes:
        return None
    if raw_bytes.startswith(_ZIP_MAGIC):
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                names = set(zf.namelist())
                if any(name.startswith("word/") for name in names):
                    return "docx"
                if "xl/workbook.xml" in names:
                    if "xl/vbaProject.bin" in names:
                        return "xlsm"
                    return "xlsx"
        except zipfile.BadZipFile:
            return None
        return None
    if raw_bytes.startswith(b"%PDF-"):
        return None
    try:
        head = raw_bytes[:4096].decode("utf-8")
    except UnicodeDecodeError:
        return None
    first_line = head.splitlines()[0] if head.splitlines() else ""
    if "," in first_line or ";" in first_line:
        return "csv"
    return None


def sanitize_file_bytes(
    raw_bytes: bytes,
    *,
    declared_name: str = "",
    declared_content_type: str = "",
    allowed_channels: tuple[str, ...] = ("csv", "xlsx", "docx"),
    max_file_bytes: int = 10_485_760,
    quality_policy: TextQualityPolicy | None = None,
) -> FileSanitizeResult:
    if len(raw_bytes) > max_file_bytes:
        return _reject(None, f"payload exceeds max_file_bytes={max_file_bytes}")

    channel = sniff_channel(raw_bytes, declared_name, declared_content_type)
    if channel is None:
        return _reject(None, "unknown or unsupported input format (allowlist reject)")
    if channel not in allowed_channels:
        return _reject(channel, f"channel {channel!r} is not allowed")
    if channel == "xlsm":
        return _reject(channel, "macro-enabled workbook rejected")

    if channel == "csv":
        result = _sanitize_csv(raw_bytes)
    elif channel == "xlsx":
        result = _sanitize_xlsx(raw_bytes)
    elif channel == "docx":
        result = _sanitize_docx(raw_bytes)
    else:
        return _reject(channel, f"unhandled channel: {channel}")

    if not result.accepted:
        return result

    policy = quality_policy or TextQualityPolicy()
    text_for_index = _indexable_text(result)
    ok, reason = assess_text_quality(text_for_index, policy)
    if not ok:
        return _reject(channel, f"text quality rejected: {reason}")
    return result


def neutralize_csv_formula_cell(value: str) -> str:
    if value and value[0] in _CSV_FORMULA_PREFIXES:
        return "'" + value
    return value


def assess_text_quality(text: str, policy: TextQualityPolicy | None = None) -> tuple[bool, str | None]:
    policy = policy or TextQualityPolicy()
    if len(text.strip()) < policy.min_text_chars:
        return False, "empty_or_too_short"

    total = len(text)
    replacement = text.count("\ufffd")
    control = sum(1 for char in text if _is_bad_control(char))
    printable = sum(1 for char in text if _is_printable_text_char(char))
    signal = sum(1 for char in text if char.isalnum() or _is_cjk(char))

    if replacement / total > policy.max_replacement_ratio:
        return False, "replacement_char_ratio"
    if control / total > policy.max_control_ratio:
        return False, "control_char_ratio"
    if printable / total < policy.min_printable_ratio:
        return False, "printable_ratio"
    if signal / total < policy.min_signal_ratio:
        return False, "low_text_signal"
    return True, None


def _sanitize_csv(raw_bytes: bytes) -> FileSanitizeResult:
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return _reject("csv", "csv: non-utf8 payload")

    rows = [[neutralize_csv_formula_cell(cell) for cell in row] for row in csv.reader(io.StringIO(text))]
    if not rows:
        return _reject("csv", "csv: empty payload")
    header = [cell.strip() for cell in rows[0]]
    records = [dict(zip(header, row)) for row in rows[1:]]
    return FileSanitizeResult(channel="csv", accepted=True, sanitized_rows=records)


def _sanitize_xlsx(raw_bytes: bytes) -> FileSanitizeResult:
    try:
        import openpyxl
    except ImportError:
        return _reject("xlsx", "xlsx: openpyxl unavailable")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001
        return _reject("xlsx", f"xlsx: parse error: {exc}")

    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(cell) if cell is not None else "" for cell in next(rows_iter)]
    except StopIteration:
        return _reject("xlsx", "xlsx: empty sheet")

    records: list[dict[str, Any]] = []
    for row in rows_iter:
        values = [neutralize_csv_formula_cell(cell) if isinstance(cell, str) else cell for cell in row]
        records.append(dict(zip(header, values)))
    return FileSanitizeResult(channel="xlsx", accepted=True, sanitized_rows=records)


def _sanitize_docx(raw_bytes: bytes) -> FileSanitizeResult:
    try:
        from defusedxml import ElementTree as SafeET
    except ImportError:
        return _reject("docx", "docx: defusedxml unavailable")

    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            if "word/document.xml" not in zf.namelist():
                return _reject("docx", "docx: missing document.xml")
            xml_bytes = zf.read("word/document.xml")
    except zipfile.BadZipFile:
        return _reject("docx", "docx: not a valid zip container")

    try:
        try:
            root = SafeET.fromstring(xml_bytes, forbid_entities=True, forbid_external=True)
        except TypeError:
            root = SafeET.fromstring(xml_bytes)
    except Exception as exc:  # noqa: BLE001
        return _reject("docx", f"docx: rejected unsafe XML: {exc}")

    text_runs = [node.text or "" for node in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")]
    if not text_runs:
        text_runs = [node.text or "" for node in root.iter() if str(node.tag).endswith("}t")]
    return FileSanitizeResult(channel="docx", accepted=True, sanitized_text="".join(text_runs))


def _indexable_text(result: FileSanitizeResult) -> str:
    if result.sanitized_text is not None:
        return result.sanitized_text
    chunks: list[str] = []
    for row in result.sanitized_rows:
        for value in row.values():
            if value is not None:
                chunks.append(str(value))
    return "\n".join(chunks)


def _reject(channel: str | None, reason: str) -> FileSanitizeResult:
    return FileSanitizeResult(
        channel=channel,
        accepted=False,
        reject_reason=reason,
        rejected=[{"reason": reason}],
    )


def _is_bad_control(char: str) -> bool:
    return ord(char) < 32 and char not in "\n\r\t"


def _is_printable_text_char(char: str) -> bool:
    return char in "\n\r\t" or char in string.printable or char.isprintable()


def _is_cjk(char: str) -> bool:
    code = ord(char)
    return (
        0x3040 <= code <= 0x30FF
        or 0x3400 <= code <= 0x4DBF
        or 0x4E00 <= code <= 0x9FFF
        or 0xF900 <= code <= 0xFAFF
    )
